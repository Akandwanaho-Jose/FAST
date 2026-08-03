"""Extract DOCX paragraphs and tables in document order for seed-data review."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_blocks(document: Document):
    body = document.element.body
    paragraph_map = {paragraph._p: paragraph for paragraph in document.paragraphs}
    table_map = {table._tbl: table for table in document.tables}

    for child in body.iterchildren():
        if child in paragraph_map:
            yield paragraph_map[child]
        elif child in table_map:
            yield table_map[child]


def clean(value: str) -> str:
    return " ".join(value.replace("\u00a0", " ").split())


def extract(source: Path) -> dict:
    document = Document(source)
    blocks = []
    paragraph_count = 0
    table_count = 0

    for index, block in enumerate(iter_blocks(document), start=1):
        if isinstance(block, Paragraph):
            text = clean(block.text)
            if not text:
                continue
            paragraph_count += 1
            blocks.append(
                {
                    "index": index,
                    "type": "paragraph",
                    "style": block.style.name if block.style else "",
                    "text": text,
                }
            )
        elif isinstance(block, Table):
            table_count += 1
            rows = []
            row_lines = []
            for row in block.rows:
                cells = [clean(cell.text) for cell in row.cells]
                if any(cells):
                    rows.append(cells)
                    row_lines.append(
                        [
                            [
                                clean(paragraph.text)
                                for paragraph in cell.paragraphs
                                if clean(paragraph.text)
                            ]
                            for cell in row.cells
                        ]
                    )
            blocks.append(
                {
                    "index": index,
                    "type": "table",
                    "rows": rows,
                    "row_lines": row_lines,
                }
            )

    nested_tables = []
    top_level_elements = {table._tbl for table in document.tables}
    for table_element in document.element.body.xpath(".//w:tbl"):
        if table_element in top_level_elements:
            continue
        table = Table(table_element, document)
        rows = []
        for row in table.rows:
            cells = [clean(cell.text) for cell in row.cells]
            if any(cells):
                rows.append(cells)
        if rows:
            nested_tables.append(rows)

    return {
        "source": str(source),
        "paragraphs": paragraph_count,
        "tables": table_count,
        "nested_tables": nested_tables,
        "blocks": blocks,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(extract(args.source), indent=2, ensure_ascii=False),
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
